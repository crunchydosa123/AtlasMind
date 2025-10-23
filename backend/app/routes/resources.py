from fastapi import APIRouter, UploadFile, File, Form, HTTPException, Depends
from pydantic import BaseModel
from supabase import create_client
from uuid import UUID
from typing import Optional, List
import pdfplumber
import docx
import openpyxl
import os, jwt
from io import BytesIO
from dotenv import load_dotenv
from datetime import datetime
from fastapi.security import OAuth2PasswordBearer
from app.services.neo4j_service import add_resource_to_graph, link_resource_to_concept, extract_concepts, add_resource_with_concepts
load_dotenv()

router = APIRouter()

SUPABASE_URL = os.environ.get("SUPABASE_URL")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY")
SUPABASE_BUCKET = os.environ.get("SUPABASE_BUCKET", "resources")
SECRET_KEY = os.environ.get("JWT_SECRET")
ALGORITHM = os.environ.get("JWT_ALGORITHM")


supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="auth/login")


def extract_text(file: UploadFile):
    file_ext = file.filename.split(".")[-1].lower()
    content = ""

    if file_ext == "pdf":
        with pdfplumber.open(file.file) as pdf:
            content = "\n".join(page.extract_text() or "" for page in pdf.pages)

    elif file_ext == "docx":
        doc = docx.Document(file.file)
        content = "\n".join([p.text for p in doc.paragraphs])

    elif file_ext == "xlsx":
        wb = openpyxl.load_workbook(file.file)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                content += " ".join(str(cell) for cell in row if cell) + "\n"

    else:
        try:
            content = file.file.read().decode("utf-8", errors="ignore")
        except Exception:
            content = ""

    return content.strip()


class ResourceResponse(BaseModel):
    id: UUID
    created_by: UUID
    project_id: UUID
    file_name: str
    file_type: str
    file_url: Optional[str]
    parsed_text: Optional[str]
    created_at: datetime


@router.get("/project/{project_id}", response_model=List[ResourceResponse])
def get_resources_by_project(project_id: UUID):
    response = supabase.table("Resources").select("*").eq("project_id", str(project_id)).execute()
    return response.data

def get_current_user(token: str = Depends(oauth2_scheme)):
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=["HS256"])
        return {"id": payload["user_id"], "email": payload["email"]}
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid token")


@router.post("/upload")
async def upload_resource(
    project_id: UUID = Form(...),
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user)
):
    file_name = file.filename
    file_type = file.content_type
    file_ext = file_name.split(".")[-1].lower()
    created_by = current_user["id"]

    try:
        file_bytes = await file.read()

        file.file.seek(0)

        parsed_text = extract_text(file)
        concepts = extract_concepts(parsed_text)

        path_on_storage = f"{project_id}/{file_name}"
        upload_resp = supabase.storage.from_(SUPABASE_BUCKET).upload(
            path_on_storage,
            file_bytes,
        )

        if isinstance(upload_resp, dict) and "error" in upload_resp:
            raise HTTPException(status_code=500, detail=f"Storage upload failed: {upload_resp['error']['message']}")

        file_url = f"{SUPABASE_URL}/storage/v1/object/public/{SUPABASE_BUCKET}/{path_on_storage}"

        insert_data = {
            "created_by": str(created_by),
            "project_id": str(project_id),
            "file_name": file_name,
            "file_type": file_type,
            "file_url": file_url,
            "parsed_text": parsed_text,
        }

        db_resp = supabase.table("Resources").insert(insert_data).execute()
        resource_id = db_resp.data[0]['id']
        add_resource_with_concepts(resource_id, file_name, project_id, concepts, uploaded_by=created_by)
        add_resource_to_graph(resource_id, file.filename, project_id, created_by)

        return {"message": "File uploaded and parsed successfully", "data": db_resp.data[0]}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    
@router.get('/')
def get_all_resources(current_user: dict = Depends(get_current_user)):
    response = (
    supabase
    .table("Resources")
    .select("id, project_id, file_name, created_at, Projects(name)")
    .eq("created_by", current_user["id"])
    .execute()
)
    if response.data:
        return response.data
    else:
        return []
